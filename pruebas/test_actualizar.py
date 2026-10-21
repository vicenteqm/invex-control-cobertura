import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'python'))
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'.build-deps'))
import importlib.util
import tempfile
import unittest
from datetime import time
from pathlib import Path
import openpyxl

spec = importlib.util.spec_from_file_location('actualizar', Path(__file__).resolve().parents[1] / 'python' / 'actualizar.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)

class CalculationTests(unittest.TestCase):
    def test_invalid_partner_uses_valid_combined_store_key(self):
        with tempfile.TemporaryDirectory() as directory:
            base,attendance=self.make_files(directory)
            book=openpyxl.load_workbook(base);sheet=book['BASE DE TIENDAS INVEX']
            col=sheet.max_column+1;sheet.cell(1,col,'SOCIO + NUM DE TIENDA')
            sheet.cell(2,1,'|');sheet.cell(2,col,'WM 1');book.save(base);book.close()
            result=module.generate(base,attendance,'2026-10')
            store=next(s for s in result['stores'] if s['id']=='WM 1')
            self.assertEqual(store['active'],2)
            self.assertEqual(sum(store['activeSchedules'].values()),2)
        self.assertEqual(module.store_identifier('|',6248,'SAMS 6248'),'SAMS 6248')
        self.assertEqual(module.store_identifier('WM',6248,'SAMS 6248'),'WM 6248')
        self.assertEqual(module.store_identifier('|',6248,'SAMS 6249'),'| 6248')
        self.assertEqual(module.store_identifier('WN',1623,'WM 1623',{'WM 1623'}),'WM 1623')
        self.assertEqual(module.store_identifier('SAMS',6234,'WM 1623',{'SAMS 6234','WM 1623'}),'SAMS 6234')
        self.assertEqual(module.store_identifier('WN',1623,'WM 6248',{'WM 6248'}),'WN 1623')

    def test_one_required_one_active_other_profile_is_not_a_headcount_vacancy(self):
        with tempfile.TemporaryDirectory() as directory:
            base,attendance=self.make_files(directory)
            book=openpyxl.load_workbook(base);book['BASE DE TIENDAS INVEX'].cell(2,8,0);book.save(base);book.close()
            book=openpyxl.load_workbook(attendance);sheet=book['Hoja1'];sheet.delete_rows(3,2);book.save(attendance);book.close()
            store=module.generate(base,attendance,'2026-10')['stores'][0]
            self.assertEqual((store['required'],store['active'],store['gap'],store['covered']),(1,1,0,1))
            self.assertEqual((store['gapSub'],store['profileAdjustment']),(1,1))

    def test_inconsistent_authorized_total_stops_calculation(self):
        with tempfile.TemporaryDirectory() as directory:
            base,attendance=self.make_files(directory)
            book=openpyxl.load_workbook(base);sheet=book['BASE DE TIENDAS INVEX'];col=sheet.max_column+1
            sheet.cell(1,col,'TOTAL DE PLANTILLA AUTORIZA');sheet.cell(2,col,99);book.save(base);book.close()
            with self.assertRaisesRegex(ValueError,'plantilla autorizada'):module.generate(base,attendance,'2026-10')

    def make_files(self, directory, conflict=False):
        base = Path(directory) / 'base.xlsx'
        attendance = Path(directory) / 'attendance.xlsx'
        b = openpyxl.Workbook(); s = b.active; s.title = 'BASE DE TIENDAS INVEX'
        s.append(['SOCIO','NUMERO DE TIENDA/DETERMINANTE','TIENDA','REGION','CIUDAD','ESTADO','ESTATUS','PROMOTORES REQUERIDOS','SUBLIDER','SUBLIDER'])
        s.append(['WM',1,'Tienda uno','REGION 1','Ciudad','Estado','OPERANDO',1,1,'Nombre privado'])
        s.append(['WM',2,'Tienda dos','REGION 1','Ciudad','Estado','CERRADA',5,1,'Nombre'])
        b.save(base); b.close()
        b = openpyxl.Workbook(); s = b.active; s.title = 'Hoja1'
        s.append(['SOCIO','DETERMINANTE','ESTATUS','PERFIL DE PUESTO BANCO','NUM DE NOMINA','NOMBRE COMPLETO',1,'Entr','Sal',2,'Entr','Sal'])
        row=['WM',1,'ACTIVO','PROMOTOR DE CREDITO','001','Nombre privado','A',time(10,30),time(18,45),None,None,None]
        s.append(row); s.append(row)
        s.append(['WM',1,'ACTIVO','PROMOTOR DE CREDITO','002','Otra persona','F',None,None,None])
        s.append(['WM',2,'ACTIVO','SUBLIDER','003','Fuera de operación','A',None,None,None])
        s.append(['WM',1,'ACTIVO TEAM','TEAM LIDER','004','Team','A',None,None,None])
        if conflict: s.append(['WM',1,'ACTIVO','SUBLIDER','001','Nombre privado','A',None,None,None])
        b.save(attendance); b.close()
        return base, attendance

    def test_profile_gaps_duplicates_and_scope(self):
        with tempfile.TemporaryDirectory() as directory:
            base, attendance = self.make_files(directory)
            result = module.generate(base, attendance, '2026-10')
            store=result['stores'][0]
            self.assertEqual((store['required'], store['active'], store['gap'], store['covered'], store['extra']), (2,2,0,2,0))
            self.assertEqual((store['profileGap'],store['profileAdjustment']),(1,1))
            self.assertEqual(store['activeSchedules'],{'Sin horario asignado':2})
            self.assertEqual(len(result['attendance']),1)
            self.assertEqual(result['attendance'][0]['codes'],{'A':1,'F':1})
            hours=result['attendance'][0]['hours']
            self.assertEqual((hours['entries'],hours['exits'],hours['completePairs'],hours['durationMinutes']),(1,1,1,495))
            self.assertEqual(result['codeLabels']['I'],'Incapacidad')
            self.assertEqual(result['cutoff'], '2026-10-01')
            self.assertEqual(result['warnings']['Filas de personal duplicadas omitidas'],1)
            self.assertNotIn('Nombre privado',str(result))

    def test_conflicting_identity_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            base, attendance = self.make_files(directory,True)
            with self.assertRaisesRegex(ValueError,'duplicado'):
                module.generate(base, attendance, '2026-10')

    def test_excel_time_formats_and_missing_values(self):
        self.assertEqual(module.clock_minutes(time(14,39)),879)
        self.assertEqual(module.clock_minutes(0.5),720)
        self.assertEqual(module.clock_minutes('20:01:00'),1201)
        for value in (None,0,'',time(0,0),'00:00:00'):
            self.assertIsNone(module.clock_minutes(value))
        with self.assertRaises(ValueError): module.clock_minutes('25:00')

    def test_incomplete_and_reversed_hours(self):
        summary=module.empty_hours()
        module.add_hours(summary,600,None)
        module.add_hours(summary,1200,600)
        self.assertEqual((summary['entries'],summary['exits'],summary['incompletePairs'],summary['reversedPairs']),(2,1,1,1))
        self.assertEqual(summary['durationPairs'],0)

if __name__ == '__main__': unittest.main()
