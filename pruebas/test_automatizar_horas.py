import csv
import os
import sys
import tempfile
import unittest
from datetime import date,datetime,timezone
from pathlib import Path
from unittest.mock import patch
sys.path[:0]=[str(Path(__file__).resolve().parents[1]/'python'),str(Path(__file__).resolve().parents[1]/'.build-deps')]
from automatizar_horas_productivas import rango_mes,MODULOS,corregir_records,ultima_bpi,ejecutar
from horas_productivas import read_productivity
from openpyxl import Workbook,load_workbook

class AutomaticHoursTests(unittest.TestCase):
    def test_full_month_mexico_leap_year_and_year_boundary(self):
        self.assertEqual(rango_mes(datetime(2026,10,15,18,tzinfo=timezone.utc)),(date(2026,10,1),date(2026,10,31)))
        self.assertEqual(rango_mes(datetime(2027,1,1,2,tzinfo=timezone.utc)),(date(2026,12,1),date(2026,12,31)))
        self.assertEqual(rango_mes(datetime(2028,2,10,18,tzinfo=timezone.utc)),(date(2028,2,1),date(2028,2,29)))

    def test_six_corrections_preserve_all_other_records(self):
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder);source=folder/'records.csv';target=folder/'corrected.csv'
            headers=['Folio FICO','F. Creación','Lugar de venta','Segundo lugar de venta','Aprobación','Estatus del proceso']
            rows=[[str(i),'2026-10-01 10:00:00','  '+k.lower()+'  ','-','Aprobada','Finalizado'] for i,k in enumerate(MODULOS)]
            rows.append(['999','2026-10-01 10:00:00','WALMART','777','Pendiente','Captura'])
            with source.open('w',encoding='utf-8-sig',newline='') as stream:
                w=csv.writer(stream);w.writerow(headers);w.writerows(rows)
            original=source.read_bytes();result=corregir_records(source,target)
            with target.open(encoding='utf-8-sig',newline='') as stream:updated=list(csv.DictReader(stream))
            self.assertEqual(result['correcciones'],6);self.assertEqual(len(updated),7)
            self.assertEqual([r['Segundo lugar de venta'] for r in updated],list(MODULOS.values())+['777'])
            self.assertEqual(source.read_bytes(),original)

    def test_latest_bpi_ignores_excel_lock(self):
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder)
            with self.assertRaises(ValueError):ultima_bpi(folder)
            for i,name in enumerate(['old.xlsx','new.xlsx','~$new.xlsx']):
                path=folder/name;path.write_bytes(b'x');os.utime(path,(100+i,100+i))
            self.assertEqual(ultima_bpi(folder).name,'new.xlsx')

    def test_download_to_monthly_report_and_failure_keeps_previous(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);bpi=root/'datos/fuentes/bpi/BPI.xlsx';bpi.parent.mkdir(parents=True)
            book=Workbook();sheet=book.active
            sheet.append(['FOLIO','EJECUTIVO P','ESTATUS BANCO','DETERMINANTE','SEDE'])
            sheet.append(['1','TEST','OK','1234','WM PRUEBA'])
            sheet.append(['2','TEST','OK','5678','SAMS PRUEBA'])
            sheet=book.create_sheet('PROIND');sheet.append(['NOMBRE','PROMOTOR']);sheet.append(['TEST','Prueba'])
            book.save(bpi);book.close()
            records=root/'datos/descargas/records/Records.csv';records.parent.mkdir(parents=True)
            with records.open('w',encoding='utf-8-sig',newline='') as stream:
                writer=csv.writer(stream);writer.writerow(['Folio FICO','F. Creación','Lugar de venta','Segundo lugar de venta','Aprobación','Estatus del proceso'])
                writer.writerow(['1','2026-10-01 10:00:00','WALMART','1234','Aprobada','Finalizado'])
                writer.writerow(['2','2026-10-15 15:00:00','SAMS','5678','Pendiente','Captura'])
            calls=[]
            def download(config,root,start,end):calls.append((start,end));return records
            config={'period':'2026-10'};now=datetime(2026,10,15,18,tzinfo=timezone.utc)
            final=ejecutar(config,root,now,download)
            self.assertEqual(calls,[(date(2026,10,1),date(2026,10,31))])
            self.assertEqual(final,root/'resultados/horas_productivas/2026/Horas Productivas Octubre.xlsx')
            result=read_productivity(final,'2026-10',[])
            self.assertEqual(result['total'],2);self.assertTrue(all(s['matches'] for s in result['sheets']))
            previous=final.read_bytes()
            def failure(*args):raise RuntimeError('Descarga fallida')
            with self.assertRaisesRegex(RuntimeError,'Descarga fallida'):ejecutar(config,root,now,failure)
            self.assertEqual(final.read_bytes(),previous)
            original_replace=Path.replace
            def locked_replace(path,target):
                if Path(target)==final:raise PermissionError('Excel abierto')
                return original_replace(path,target)
            with patch.object(Path,'replace',locked_replace):
                alternative=ejecutar(config,root,now,download)
            self.assertNotEqual(alternative,final)
            self.assertTrue(alternative.is_file())
            self.assertEqual(final.read_bytes(),previous)
            self.assertEqual(read_productivity(alternative,'2026-10',[])['total'],2)
            with self.assertRaises(ValueError):ejecutar({'period':'2026-09'},root,now,download)
