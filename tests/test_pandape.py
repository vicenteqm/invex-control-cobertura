import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'.build-deps'))
import tempfile
import unittest
from datetime import date
from openpyxl import Workbook,load_workbook
from generar_asistencia import generate,summarize,attendance_code

class PandapeTests(unittest.TestCase):
    def fixtures(self,folder):
        report=folder/'Reportede.xlsx';w=Workbook();s=w.active
        s.cell(5,1,'Fecha inicial:');s.cell(5,2,'13/10/2026');s.cell(6,1,'Fecha final:');s.cell(6,2,'13/10/2026')
        s.cell(11,3,'Colaborador')
        for column in (13,24,35):s.cell(11,column,'Ent.')
        for column in (18,29,40):s.cell(11,column,'Sal.')
        s.cell(12,1,'13/10/2026');s.cell(13,1,'13/10/2026')
        s.cell(12,3,'Pérez Ana');s.cell(12,13,'10:12');s.cell(12,18,'17:00')
        s.cell(13,3,'Ana Perez');s.cell(13,24,'10:05');s.cell(13,29,'18:00');w.save(report);w.close()
        source=folder/'source.xlsx';w=Workbook();s=w.active;s.title='Hoja1'
        s.append(['NOMBRE COMPLETO','ESTATUS','DESCANSO OFICIAL',13,'Entr','Sal'])
        s.append(['Ana Pérez','ACTIVO','MARTES',None,None,None]);s.append(['Otra persona','ACTIVO','LUNES','V',None,None]);w.save(source);w.close()
        return source,report
    def test_consolidation_and_idempotent_column(self):
        with tempfile.TemporaryDirectory() as d:
            folder=Path(d);source,report=self.fixtures(folder);target=folder/'output.xlsx'
            result=generate(source,report,target,date(2026,10,13))
            self.assertEqual(result['matched'],1);self.assertEqual(result['protected'],1)
            generate(target,report,target,date(2026,10,13))
            w=load_workbook(target);s=w['Hoja1']
            self.assertEqual(s.max_column,7);self.assertEqual(s.cell(2,4).value,'R');self.assertEqual(s.cell(3,4).value,'V')
            self.assertAlmostEqual(s.cell(2,5).value.hour*60+s.cell(2,5).value.minute,605);w.close()
    def test_original_rest_and_tardy_rules(self):
        self.assertEqual(attendance_code(None,None,'MARTES',date(2026,10,13)),'D')
        self.assertEqual(attendance_code(None,None,'LUNES',date(2026,10,13)),'F')
        self.assertEqual(attendance_code(616,None,'',date(2026,10,13)),'A')
    def test_report_date_must_match_capture_day(self):
        with tempfile.TemporaryDirectory() as d:
            _,report=self.fixtures(Path(d))
            with self.assertRaises(ValueError):summarize(report,date(2026,10,14))
    def test_range_filters_days_and_reads_fourth_pair(self):
        with tempfile.TemporaryDirectory() as d:
            _,report=self.fixtures(Path(d))
            w=load_workbook(report);s=w.active;s.cell(5,2,'13/09/2026')
            s.cell(14,1,'12/10/2026');s.cell(14,3,'Ana Perez');s.cell(14,13,'07:00');s.cell(14,18,'23:00')
            s.cell(11,46,'Ent.   ');s.cell(11,51,'Sal.   ');s.cell(12,46,'10:02');s.cell(12,51,'19:00');w.save(report);w.close()
            values=list(summarize(report,date(2026,10,13)).values())
            self.assertEqual(len(values),1);self.assertEqual(values[0]['entry'],602);self.assertEqual(values[0]['exit'],1140)
    def test_daily_report_with_two_pairs(self):
        with tempfile.TemporaryDirectory() as d:
            _,report=self.fixtures(Path(d))
            w=load_workbook(report);w.active.delete_cols(35,6);w.save(report);w.close()
            values=list(summarize(report,date(2026,10,13)).values())
            self.assertEqual(len(values),1);self.assertEqual(values[0]['entry'],605);self.assertEqual(values[0]['exit'],1080)
    def test_unmatched_report_preserves_previous_file(self):
        with tempfile.TemporaryDirectory() as d:
            folder=Path(d);source,report=self.fixtures(folder);target=folder/'output.xlsx';target.write_bytes(b'previo')
            w=load_workbook(report);w.active.cell(12,3,'Desconocido');w.active.cell(13,3,'Desconocido');w.save(report);w.close()
            with self.assertRaises(ValueError):generate(source,report,target,date(2026,10,13))
            self.assertEqual(target.read_bytes(),b'previo')
