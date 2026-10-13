"""Adaptador del motor original al panel INVEX y a encabezados de Excel."""
import calendar
import json
from datetime import datetime, timezone, timedelta
from pathlib import Path
import openpyxl
import engine
from actualizar import norm, identifier
from importar_operacion import sanitize_months

def local_now():
    try:
        from zoneinfo import ZoneInfo
        return datetime.now(ZoneInfo('America/Mexico_City'))
    except (ImportError, KeyError):
        return datetime.now(timezone(timedelta(hours=-6)))

def read_rows(path, sheet):
    book=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        if sheet not in book.sheetnames: raise ValueError('Falta la hoja '+sheet)
        return [list(row) for row in book[sheet].iter_rows(values_only=True)]
    finally: book.close()

def adapt_rows(rows, kind, days):
    heads={}
    for i,h in enumerate(rows[0]):heads.setdefault(norm(h),[]).append(i)
    mappings = {1:'REGION',2:'ESTADO',3:'CIUDAD',4:'SOCIO',5:'DETERMINANTE',7:'NOMBRE DE LA TIENDA',8:'PERFIL DE PUESTO BANCO',9:'NOMBRE COMPLETO',10:'ESTATUS',11:'USUARIO',13:'PIN PANDAPE',14:'JORNADA',15:'HORARIO',16:'FECHA DE ALTA',17:'FECHA DE INGRESO A PDV',18:'FECHA DE BAJA',20:'DESCANSO OFICIAL'} if kind=='attendance' else {0:'PROYECTO',1:'ZONAS',3:'REGION',4:'SUB REGION OPERATIVA',5:'ESTADO',6:'CIUDAD',7:'SOCIO',8:'NUMERO DE TIENDA/DETERMINANTE',10:'TIENDA',19:'ESTATUS',20:'HORARIO (RESPETAR FORMATO DE HORA)',25:'TOTAL DE PLANTILLA AUTORIZA',36:'EJECUTIVA RH'}
    mandatory=['SOCIO','DETERMINANTE','NOMBRE COMPLETO','ESTATUS','USUARIO','HORARIO','FECHA DE INGRESO A PDV','FECHA DE BAJA','DESCANSO OFICIAL'] if kind=='attendance' else ['SOCIO','NUMERO DE TIENDA/DETERMINANTE','TIENDA','ESTATUS','TOTAL DE PLANTILLA AUTORIZA']
    missing=[h for h in mandatory if h not in heads]
    if missing:raise ValueError('Faltan columnas del motor: '+', '.join(missing))
    header=[None]*max(47,22+days*3)
    output=[header];seen={}
    for original in rows[1:]:
        row=[None]*len(header)
        for index,name in mappings.items():
            if name in heads:row[index]=original[heads[name][-1 if name=='REGION' and kind=='attendance' else 0]]
        key_name='SOCIO Y NUMERO DE TIENDA' if kind=='attendance' else 'SOCIO + NUM DE TIENDA'
        cached=norm(original[heads[key_name][0]]) if key_name in heads else ''
        socio=row[4 if kind=='attendance' else 7];det=row[5 if kind=='attendance' else 8]
        if not det:continue
        key=cached or norm(socio)+' '+identifier(det)
        row[6 if kind=='attendance' else 9]=key
        if kind=='attendance':
            row[10]=norm(row[10]);row[20]=norm(row[20])
            for day in range(1,days+1):
                candidates=heads.get(str(day),[])
                if not candidates:continue
                src=candidates[0]
                if src+2>=len(original) or norm(rows[0][src+1])!='ENTR' or norm(rows[0][src+2])!='SAL':raise ValueError('Día sin Entr / Sal')
                dst=22+(day-1)*3;row[dst:dst+3]=[norm(original[src]) or None,original[src+1],original[src+2]]
            if not row[9]:continue
            payroll=norm(original[heads['NUM DE NOMINA'][0]]) if 'NUM DE NOMINA' in heads else ''
            identity=(payroll if payroll!='PENDIENTE' else '',norm(row[9]))
            fingerprint=tuple(str(v) for v in row)
            if row[10]=='ACTIVO':
                if identity in seen:
                    if seen[identity]!=fingerprint:raise ValueError('Personal activo duplicado con datos distintos')
                    continue
                seen[identity]=fingerprint
        else:
            if key in seen:raise ValueError('Tienda duplicada: '+key)
            seen[key]=True
            if len(heads.get('SUBLIDER',[]))>1:row[32]=original[heads['SUBLIDER'][1]]
            if norm(row[32]) in ('PENDIENTE','N/A','SIN ASIGNAR','SIN SUBLIDER'):row[32]=None
        output.append(row)
    return output

def process(base, attendance, period, maat=None, now=None):
    year,month=map(int,period.split('-'));days=calendar.monthrange(year,month)[1]
    a=adapt_rows(read_rows(attendance,'Hoja1'),'attendance',days)
    b=adapt_rows(read_rows(base,'BASE DE TIENDAS INVEX'),'stores',days)
    employees=engine.build_employees(a,days)
    reference=max((d for employee in employees if employee['estatus']=='ACTIVO' for d in employee['dias']),default=0)
    if not reference:raise ValueError('Sin captura para calcular cobertura operativa')
    now=now or local_now()
    cutoff=(now.hour*60+now.minute) if (now.year,now.month,now.day)==(year,month,reference) else None
    kwargs={};available=set()
    if maat:
        from descarga_asistencia import read_csv
        headers,records=read_csv(maat)
        rows=[headers]+[[r.get(h,'') for h in headers] for r in records]
        values=engine.load_maat_usuarios(rows,anio_esperado=year,mes_esperado=month)
        names=['maat_usuarios_por_dia','maat_ubicacion_por_usuario_dia','maat_entrada_por_usuario_dia','maat_salida_por_usuario_dia','maat_cierre_manual_por_usuario_dia','maat_dia_min','maat_dia_max']
        kwargs=dict(zip(names,values));available=set(values[0])
    raw=engine.analyze(a,b,days,reference_day_override=reference,hora_corte_min=cutoff,anio=year,mes=month,**kwargs)
    private=Path(__file__).resolve().parent/'inputs'
    private.mkdir(exist_ok=True)
    (private/'analisis_actual.json').write_text(json.dumps(raw,ensure_ascii=False,default=str),encoding='utf-8')
    # Los días sin captura global permanecen sin datos, no se convierten en alarmas.
    captured={d for employee in employees if employee['estatus']=='ACTIVO' for d in employee['dias']}
    for store in raw['stores'].values():
        store['dias']={d:r for d,r in store['dias'].items() if int(d) in captured}
        observed=list(store['dias'].values())
        store['dias_rojo_mes']=sum(r['semaforo'] in ('ALARMANTE','CRITICA') for r in observed)
        store['promedio_cobertura_mes']=round(sum(r['cobertura_pct'] for r in observed)/len(observed),1) if observed else 0
    result=sanitize_months({period:raw})[period]
    result.update(mode='current',processedAt=now.isoformat(),cutoffTime=engine.fmt_min(cutoff) if cutoff is not None else 'Corte histórico del Excel',maatLoaded=bool(maat),maatDays=sorted(available),attendanceDays=sorted(captured))
    aggregates={s['id']:s for s in result['stores']}
    for store in aggregates.values():
        store['registration']={str(day):dict(available=day in available,pandape=0,maat=0,both=0,none=0,movedOut=0,incoming=0,late=0,earlyDeparture=0,manualClosing=0) for day in range(1,days+1)}
    for key,store in raw['stores'].items():
        for employee in store['empleados_activos']+store.get('externos_maat',[]):
            for day_str,record in employee.get('registro_por_dia',{}).items():
                target=record.get('tienda_maat_key') if record.get('se_movio_de_tienda') else key
                if target not in aggregates:continue
                bucket=aggregates[target]['registration'][day_str]
                pandape=bool(record.get('registra_pandape'));maat_record=bool(record.get('registra_maat'))
                bucket['pandape']+=pandape;bucket['maat']+=maat_record;bucket['both']+=pandape and maat_record
                bucket['none']+=not pandape and not maat_record
                bucket['incoming']+=target!=key
                for field in ['late','earlyDeparture','manualClosing']:
                    source={'late':'llego_tarde','earlyDeparture':'salio_antes_de_horario','manualClosing':'cerro_jornada_manual'}[field]
                    bucket[field]+=bool(record.get(source))
                if target!=key and key in aggregates:aggregates[key]['registration'][day_str]['movedOut']+=1
    return result
