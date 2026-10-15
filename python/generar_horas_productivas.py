"""Reporte por tienda, día y hora para Walmart y Sam's, con estados de Records."""
from pathlib import Path
import re
import unicodedata
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ESTADOS = ['Aprobada con Garantía', 'Pendiente', 'Rechazada', 'Aprobada']
OTROS = 'Otras aprobaciones'
FINALIZADAS = 'Aprobadas finalizadas'
METRICAS = ESTADOS + [OTROS, FINALIZADAS]
CLAVES = ['CADENA', 'DETERMINANTE', 'SEDE']


def normalizar(valor):
    if pd.isna(valor):
        return ''
    texto = ''.join(c for c in unicodedata.normalize('NFD', str(valor)) if not unicodedata.combining(c))
    return re.sub(r'\s+', ' ', texto).strip().upper()


def folios(serie):
    return serie.astype('string').str.strip().str.replace(r'\.0$', '', regex=True).replace(
        {'': pd.NA, '-': pd.NA, 'nan': pd.NA, 'None': pd.NA, '<NA>': pd.NA})


def cadena(sede):
    texto = normalizar(sede).replace("'", '').replace('’', '')
    if re.match(r'^(WM|WALMART)(?:\b|\d)', texto):
        return 'Walmart'
    if re.match(r'^SAMS?(?:\b|\d)', texto):
        return "Sam's"
    return ''


def leer_records(archivo):
    if Path(archivo).suffix.lower() in ('.xlsx', '.xls'):
        tabla = pd.read_excel(archivo, dtype=str)
    else:
        tabla = pd.read_csv(archivo, dtype=str, encoding='utf-8-sig', sep=None, engine='python')
    tabla.columns = tabla.columns.str.strip()
    # Localizar encabezados sin depender de mayúsculas o acentos.
    esperadas = ['Folio FICO', 'F. Creación', 'Aprobación', 'Estatus del proceso']
    disponibles = {normalizar(c): c for c in tabla.columns}
    faltantes = [c for c in esperadas if normalizar(c) not in disponibles]
    if faltantes:
        raise ValueError('Faltan columnas en Records: ' + ', '.join(faltantes))
    return tabla.rename(columns={disponibles[normalizar(c)]: c for c in esperadas})


def preparar_datos(bpi, records, inicio=None, fin=None):
    print('Cargando BPI y Records...')
    base = pd.read_excel(bpi, dtype=str)
    proind = pd.read_excel(bpi, sheet_name='PROIND', dtype=str)
    base.columns = base.columns.str.strip()
    proind.columns = proind.columns.str.strip()
    for tabla, requeridas, nombre in [
        (base, ['FOLIO', 'EJECUTIVO P', 'DETERMINANTE', 'SEDE', 'ESTATUS BANCO'], 'BPI'),
        (proind, ['NOMBRE', 'PROMOTOR'], 'PROIND'),
    ]:
        faltantes = set(requeridas) - set(tabla.columns)
        if faltantes:
            raise ValueError(f'Faltan columnas en {nombre}: {", ".join(sorted(faltantes))}')
    base['FOLIO'] = folios(base['FOLIO'])
    base = base[base['FOLIO'].notna()].copy()
    for col in ['EJECUTIVO P', 'DETERMINANTE', 'SEDE']:
        base[col] = base[col].fillna('').str.strip()
    base['DETERMINANTE'] = base['DETERMINANTE'].str.replace(r'\.0$', '', regex=True)
    repetidos = base[base['FOLIO'].duplicated(False)]
    if not repetidos.empty and repetidos.groupby('FOLIO')[['DETERMINANTE', 'SEDE', 'EJECUTIVO P']].nunique().gt(1).any().any():
        raise ValueError('El BPI asigna el mismo folio a distintas tiendas o ejecutivos. Revisa los duplicados.')
    base = base.drop_duplicates('FOLIO')
    base['CADENA'] = base['SEDE'].map(cadena)
    base = base[base['CADENA'] != ''].copy()

    registros = leer_records(records)
    registros['FOLIO'] = folios(registros['Folio FICO'])
    registros['FECHA HORA CREACION'] = pd.to_datetime(registros['F. Creación'], errors='coerce', format='mixed')
    registros = registros[registros['FOLIO'].notna()].copy()
    # Una fila por folio: primera fecha de creación y estado de la aparición más reciente.
    fechas = registros.groupby('FOLIO')['FECHA HORA CREACION'].min()
    registros = registros.sort_values('FECHA HORA CREACION', kind='stable', na_position='first').drop_duplicates('FOLIO', keep='last')
    registros['FECHA HORA CREACION'] = registros['FOLIO'].map(fechas)
    registros = registros[registros['FECHA HORA CREACION'].notna()].copy()
    if inicio is not None:
        registros = registros[registros['FECHA HORA CREACION'] >= pd.Timestamp(inicio)]
    if fin is not None:
        registros = registros[registros['FECHA HORA CREACION'] < pd.Timestamp(fin) + pd.Timedelta(days=1)]
    datos = base.merge(registros[['FOLIO', 'FECHA HORA CREACION', 'Aprobación', 'Estatus del proceso']], on='FOLIO', how='inner', validate='one_to_one')
    print(f'Tiendas Walmart y Sam\'s en BPI: {base[CLAVES].drop_duplicates().shape[0]}')
    print(f'Folios BPI con fecha de Records en el periodo: {len(datos)}; sin coincidencia en el periodo: {len(base) - len(datos)}')
    datos['FECHA'] = datos['FECHA HORA CREACION'].dt.strftime('%Y-%m-%d')
    datos['RANGO HORA'] = datos['FECHA HORA CREACION'].dt.hour.map(lambda h: f'{h:02d}:00 - {h:02d}:59')
    proind['NOMBRE'] = proind['NOMBRE'].fillna('').str.strip()
    nombres = proind.drop_duplicates('NOMBRE').set_index('NOMBRE')['PROMOTOR']
    datos['NOMBRE_COMPLETO'] = datos['EJECUTIVO P'].map(nombres).fillna('No encontrado')
    datos['Aprobación'] = datos['Aprobación'].fillna('Sin dato').str.strip().replace({'': 'Sin dato', '-': 'Sin dato'})
    datos['Estatus del proceso'] = datos['Estatus del proceso'].fillna('Sin dato').str.strip().replace({'': 'Sin dato', '-': 'Sin dato'})
    canonicos = {normalizar(e): e for e in ESTADOS}
    datos['CATEGORIA'] = datos['Aprobación'].map(normalizar).map(canonicos).fillna(OTROS)
    for estado in ESTADOS + [OTROS]:
        datos[estado] = datos['CATEGORIA'].eq(estado).astype(int)
    datos[FINALIZADAS] = (datos['CATEGORIA'].isin(['Aprobada', 'Aprobada con Garantía']) &
        datos['Estatus del proceso'].map(normalizar).eq('FINALIZADO')).astype(int)
    return datos


def resumen(datos, claves):
    conteos = datos.groupby(claves, dropna=False).size().rename('TOTAL_TRAMITES')
    metricas = datos.groupby(claves, dropna=False)[METRICAS].sum()
    return conteos.to_frame().join(metricas).reset_index()


def tabla_tiendas(datos, tiendas, horas):
    columnas = CLAVES + horas + ['TOTAL_TRAMITES'] + METRICAS
    if datos.empty:
        tabla = tiendas.copy()
        for col in horas + ['TOTAL_TRAMITES'] + METRICAS:
            tabla[col] = 0
    else:
        grid = datos.groupby(CLAVES + ['RANGO HORA']).size().unstack('RANGO HORA', fill_value=0).reindex(columns=horas, fill_value=0).reset_index()
        tabla = tiendas.merge(grid, on=CLAVES, how='left').merge(resumen(datos, CLAVES), on=CLAVES, how='left')
        tabla[horas + ['TOTAL_TRAMITES'] + METRICAS] = tabla[horas + ['TOTAL_TRAMITES'] + METRICAS].fillna(0).astype(int)
    tabla = tabla[columnas].sort_values(['CADENA', 'TOTAL_TRAMITES', 'SEDE'], ascending=[True, False, True], kind='stable')
    total = {col: '' for col in columnas}
    total['DETERMINANTE'] = 'TOTAL'
    for col in horas + ['TOTAL_TRAMITES'] + METRICAS:
        total[col] = int(tabla[col].sum())
    return pd.concat([tabla, pd.DataFrame([total])], ignore_index=True)


def hora_pico(datos, claves):
    por_hora = datos.groupby(claves + ['RANGO HORA']).size().rename('TRAMITES EN HORA PICO').reset_index()
    pico = por_hora.sort_values(claves + ['TRAMITES EN HORA PICO', 'RANGO HORA'], ascending=[True]*len(claves) + [False, True]).drop_duplicates(claves)
    pico = pico.rename(columns={'RANGO HORA': 'HORA PICO'}).merge(resumen(datos, claves), on=claves)
    pico['% EN HORA PICO'] = (pico['TRAMITES EN HORA PICO'] / pico['TOTAL_TRAMITES']).round(4)
    return pico


def formato_tabla(hoja, fila, tabla, total=False):
    cabecera = PatternFill('solid', fgColor='000000')
    for celda in hoja[fila]:
        if celda.column <= len(tabla.columns):
            celda.font = Font(name='Calibri', size=9, bold=True, color='FFFFFF')
            celda.fill = cabecera
            celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    hoja.row_dimensions[fila].height = 44
    for j, nombre in enumerate(tabla.columns, 1):
        for row in range(fila+1, fila+len(tabla)+1):
            celda = hoja.cell(row, j)
            celda.font = Font(name='Calibri', size=9, color='000000')
            celda.fill = PatternFill('solid', fgColor='FFFFFF')
            celda.alignment = Alignment(vertical='center', horizontal='center')
            if '%' in nombre:
                celda.number_format = '0.0%'
            elif nombre == 'FECHA HORA CREACION':
                celda.number_format = 'dd/mm/yyyy hh:mm:ss'
            elif isinstance(celda.value, (float, int)):
                celda.number_format = '#,##0'
            # Mantener folios/identificadores y textos del origen como valores literales.
            if isinstance(celda.value, str) and celda.value.startswith('='):
                celda.data_type = 's'
    if total:
        for celda in hoja[fila+len(tabla)]:
            if celda.column <= len(tabla.columns):
                celda.font = Font(name='Calibri', size=9, bold=True)
                celda.fill = PatternFill('solid', fgColor='FFFFFF')


def escribir_reporte(datos, salida, inicio=None, fin=None):
    horas = sorted(datos['RANGO HORA'].unique())
    tiendas = datos[CLAVES].drop_duplicates()
    with pd.ExcelWriter(salida, engine='openpyxl') as writer:
        bloques = [('GENERAL', datos)] + [(dia, datos[datos['FECHA'] == dia]) for dia in sorted(datos['FECHA'].unique())]
        for nombre, periodo in bloques:
            tabla = tabla_tiendas(periodo, tiendas, horas)
            tabla.to_excel(writer, sheet_name=nombre, index=False)
            hoja = writer.sheets[nombre]
            formato_tabla(hoja, 1, tabla, total=True)
            hoja.freeze_panes = 'D2'
            hoja.auto_filter.ref = f'A1:{get_column_letter(len(tabla.columns))}{max(1, len(tabla))}'
            fila = len(tabla) + 4
            hoja.cell(fila, 1, 'APROBACIÓN Y PROCESO POR HORA').font = Font(name='Calibri', bold=True, size=9)
            por_hora = resumen(periodo, ['RANGO HORA'])
            por_hora.to_excel(writer, sheet_name=nombre, startrow=fila, index=False)
            formato_tabla(hoja, fila+1, por_hora)
            fila += len(por_hora) + 4
            hoja.cell(fila, 1, 'Aprobadas finalizadas: Aprobada o Aprobada con Garantía y proceso Finalizado. Es un subconjunto de las aprobadas.').font = Font(name='Calibri', italic=True, size=9)
            hoja.cell(fila+1, 1, 'Otras aprobaciones incluye estados distintos de las cuatro categorías, por ejemplo Rechazada por cliente o Sin dato.').font = Font(name='Calibri', italic=True, size=9)

        ranking = resumen(datos, ['RANGO HORA']).sort_values('TOTAL_TRAMITES', ascending=False, kind='stable').reset_index(drop=True)
        ranking.insert(0, 'LUGAR', range(1, len(ranking)+1))
        ranking['% DEL TOTAL'] = (ranking['TOTAL_TRAMITES'] / len(datos)).round(4) if len(datos) else pd.Series(dtype=float)
        ranking['% ACUMULADO'] = ranking['% DEL TOTAL'].cumsum()
        fila = 1
        for titulo, tabla in [
            ('RANKING GENERAL DE HORAS DEL PERIODO', ranking),
            ('HORA PICO DE CADA DIA', hora_pico(datos, ['FECHA'])),
            ('HORA PICO DE CADA TIENDA', hora_pico(datos, CLAVES)),
        ]:
            tabla.to_excel(writer, sheet_name='HORAS PICO', startrow=fila, index=False)
            hoja = writer.sheets['HORAS PICO']
            hoja.cell(fila, 1, titulo).font = Font(name='Calibri', size=9, bold=True)
            formato_tabla(hoja, fila+1, tabla)
            fila += len(tabla) + 4

        columnas = CLAVES + ['FOLIO', 'FECHA HORA CREACION', 'RANGO HORA', 'EJECUTIVO P',
            'NOMBRE_COMPLETO', 'ESTATUS BANCO', 'Aprobación', 'Estatus del proceso', FINALIZADAS]
        detalle = datos[columnas].sort_values(['FECHA HORA CREACION', 'CADENA', 'SEDE'])
        detalle.to_excel(writer, sheet_name='DETALLE', index=False)
        hoja = writer.sheets['DETALLE']
        formato_tabla(hoja, 1, detalle)
        hoja.freeze_panes = 'D2'
        hoja.auto_filter.ref = hoja.dimensions
        fila = len(detalle) + 4
        hoja.cell(fila, 1, 'Población: folios del BPI de Walmart y Sam\'s con fecha de creación de Records dentro del periodo. Un trámite por folio.').font = Font(name='Calibri', italic=True, size=9)
        hoja.cell(fila+1, 1, 'Duplicados Records: primera fecha de creación y estado de la aparición con creación más reciente.').font = Font(name='Calibri', italic=True, size=9)
        if inicio and fin:
            hoja.cell(fila+2, 1, f'Periodo incluido: {inicio:%d/%m/%Y} al {fin:%d/%m/%Y}.')
        for hoja in writer.sheets.values():
            for fila_celdas in hoja:
                for celda in fila_celdas:
                    celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=celda.alignment.wrap_text)
            # Anchos compartidos por columna: contemplar cada tabla de la misma hoja.
            for j in range(1, hoja.max_column+1):
                valores = [str(hoja.cell(i,j).value or '') for i in range(1, hoja.max_row+1)
                    if len(str(hoja.cell(i,j).value or '')) < 65]
                ancho = min(max(max(map(len, valores), default=8) + 2, 15), 40)
                hoja.column_dimensions[get_column_letter(j)].width = ancho
    print(f'Reporte generado: {len(datos)} trámites en {len(tiendas)} tiendas Walmart y Sam\'s.')


def generar(bpi, records, salida, inicio=None, fin=None):
    datos = preparar_datos(bpi, records, inicio, fin)
    escribir_reporte(datos, salida, inicio, fin)
