"""Genera datos agregados para INVEX. Los Excel permanecen en el equipo."""
import argparse
import json
import math
import subprocess
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone, time
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CODE_LABELS = {'A': 'Asistencia', 'F': 'Falta', 'R': 'Retardo', 'V': 'Vacaciones', 'I': 'Incapacidad'}

def clock_minutes(value):
    """Excel usa 00:00:00 como marcador sin registro, no como una checada."""
    if value is None or value == '' or value == 0:
        return None
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        minutes = value.hour * 60 + value.minute + value.second / 60
    elif isinstance(value, (int, float)) and not isinstance(value, bool) and 0 < value < 1:
        minutes = float(value) * 1440
    elif isinstance(value, str):
        try:
            parsed = time.fromisoformat(value.strip())
            minutes = parsed.hour * 60 + parsed.minute + parsed.second / 60
        except ValueError:
            raise ValueError('Formato de hora inválido') from None
    else:
        raise ValueError('Formato de hora inválido')
    return round(minutes, 4) if minutes > 0 else None

def empty_hours():
    return dict(entries=0, exits=0, completePairs=0, durationPairs=0, durationMinutes=0, incompletePairs=0, reversedPairs=0, entryFirst=None, entryLast=None, exitFirst=None, exitLast=None)

def add_hours(summary, entry, exit):
    for name, value in [('entry', entry), ('exit', exit)]:
        if value is not None:
            summary['entries' if name == 'entry' else 'exits'] += 1
            summary[name + 'First'] = min(summary[name + 'First'], value) if summary[name + 'First'] is not None else value
            summary[name + 'Last'] = max(summary[name + 'Last'], value) if summary[name + 'Last'] is not None else value
    if entry is not None and exit is not None:
        summary['completePairs'] += 1
        if exit >= entry:
            summary['durationPairs'] += 1
            summary['durationMinutes'] += exit - entry
        else:
            summary['reversedPairs'] += 1
    elif entry is not None or exit is not None:
        summary['incompletePairs'] += 1

def norm(value):
    text = str(value or '').strip().upper()
    return ' '.join(''.join(c for c in unicodedata.normalize('NFKD', text) if not unicodedata.combining(c)).split())

def identifier(value):
    if isinstance(value, (float, int)):
        return str(int(value)) if float(value).is_integer() else str(value)
    return norm(value)

def columns(sheet):
    headers = next(sheet.iter_rows(values_only=True))
    result = {}
    for i, h in enumerate(headers):
        result.setdefault(norm(h), i)
    return headers, result

def required(headers, names):
    missing = [name for name in names if name not in headers]
    if missing:
        raise ValueError('Faltan columnas: ' + ', '.join(missing))

def quantity(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0 or int(value) != value:
        raise ValueError(f'Plantilla inválida: {value!r}')
    return int(value)

def generate(base, attendance, period):
    datetime.strptime(period, '%Y-%m')
    warnings = Counter()
    book = openpyxl.load_workbook(base, read_only=True, data_only=True)
    try:
        sheet = book['BASE DE TIENDAS INVEX']
        _, h = columns(sheet)
        required(h, ['SOCIO', 'NUMERO DE TIENDA/DETERMINANTE', 'TIENDA', 'REGION', 'CIUDAD', 'ESTADO', 'ESTATUS', 'PROMOTORES REQUERIDOS', 'SUBLIDER'])
        stores = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            get = lambda name: row[h[name]]
            if not get('NUMERO DE TIENDA/DETERMINANTE'):
                continue
            key = norm(get('SOCIO')) + ' ' + identifier(get('NUMERO DE TIENDA/DETERMINANTE'))
            if key in stores:
                raise ValueError('Tienda duplicada en base: ' + key)
            operating = norm(get('ESTATUS')) == 'OPERANDO'
            stores[key] = dict(id=key, name=str(get('TIENDA') or ''), region=str(get('REGION') or 'Sin región'), city=str(get('CIUDAD') or ''), state=str(get('ESTADO') or ''), status=norm(get('ESTATUS')), requiredProm=quantity(get('PROMOTORES REQUERIDOS')) if operating else 0, requiredSub=quantity(get('SUBLIDER')) if operating else 0, activeProm=0, activeSub=0)
    finally:
        book.close()
    book = openpyxl.load_workbook(attendance, read_only=True, data_only=True)
    days = defaultdict(Counter)
    hours = defaultdict(empty_hours)
    seen = {}
    payroll_names = {}
    try:
        sheet = book['Hoja1']
        raw_headers, h = columns(sheet)
        required(h, ['SOCIO', 'DETERMINANTE', 'ESTATUS', 'PERFIL DE PUESTO BANCO', 'NUM DE NOMINA', 'NOMBRE COMPLETO'])
        day_columns = [(i, int(v)) for i, v in enumerate(raw_headers) if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() and 1 <= v <= 31]
        if not day_columns:
            raise ValueError('No se encontraron columnas de días')
        for i, day in day_columns:
            if i + 2 >= len(raw_headers) or norm(raw_headers[i + 1]) != 'ENTR' or norm(raw_headers[i + 2]) != 'SAL':
                raise ValueError(f'El día {day} debe estar seguido de Entr y Sal')
        import calendar
        max_day = calendar.monthrange(*map(int, period.split('-')))[1]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            get = lambda name: row[h[name]]
            if norm(get('ESTATUS')) != 'ACTIVO':
                continue
            role = norm(get('PERFIL DE PUESTO BANCO'))
            if role not in ('PROMOTOR DE CREDITO', 'SUBLIDER'):
                continue
            key = norm(get('SOCIO')) + ' ' + identifier(get('DETERMINANTE'))
            if key not in stores:
                warnings['Personal activo sin tienda en base'] += 1
                continue
            if stores[key]['status'] != 'OPERANDO':
                warnings['Personal activo en tienda fuera de operación'] += 1
                continue
            payroll = norm(get('NUM DE NOMINA'))
            name = norm(get('NOMBRE COMPLETO'))
            if payroll and payroll != 'PENDIENTE':
                if payroll in payroll_names and payroll_names[payroll] != name:
                    warnings['Nóminas compartidas entre nombres distintos; revisar fuente'] += 1
                payroll_names[payroll] = name
                employee = payroll + '|' + name
            else:
                employee = name
            if not employee:
                raise ValueError('Personal elegible sin identificador')
            fingerprint = (key, role, tuple(tuple(norm(row[j]) for j in (i, i + 1, i + 2)) for i, d in day_columns))
            if employee in seen:
                if seen[employee] != fingerprint:
                    raise ValueError('Identificador de personal duplicado con datos distintos; revise el Excel')
                warnings['Filas de personal duplicadas omitidas'] += 1
                continue
            seen[employee] = fingerprint
            stores[key]['activeProm' if role == 'PROMOTOR DE CREDITO' else 'activeSub'] += 1
            schedule = str(row[h['HORARIO']] or '').strip() if 'HORARIO' in h else ''
            schedule = ' '.join(schedule.split()) or 'Sin horario asignado'
            schedules = stores[key].setdefault('activeSchedules', {})
            schedules[schedule] = schedules.get(schedule, 0) + 1
            for i, day in day_columns:
                code = norm(row[i])
                if day > max_day:
                    continue
                if code:
                    days[day][code] += 1
                parsed = []
                for j in (i + 1, i + 2):
                    try:
                        parsed.append(clock_minutes(row[j]))
                    except ValueError:
                        warnings['Horas inválidas omitidas; revisar fuente'] += 1
                        parsed.append(None)
                if code or any(value is not None for value in parsed):
                    days[day]  # También conserva días con horas y sin código.
                    add_hours(hours[day], *parsed)
    finally:
        book.close()
    for store in stores.values():
        store['required'] = store['requiredProm'] + store['requiredSub']
        store['active'] = store['activeProm'] + store['activeSub']
        store['gapProm'] = max(0, store['requiredProm'] - store['activeProm'])
        store['gapSub'] = max(0, store['requiredSub'] - store['activeSub'])
        store['gap'] = store['gapProm'] + store['gapSub']
        store['covered'] = min(store['activeProm'], store['requiredProm']) + min(store['activeSub'], store['requiredSub'])
        store['extra'] = store['active'] - store['covered']
    if sum(summary['reversedPairs'] for summary in hours.values()):
        warnings['Salidas anteriores a entrada; excluidas de duración promedio'] = sum(summary['reversedPairs'] for summary in hours.values())
    code_days = [day for day, codes in days.items() if codes]
    from turnos_cobertura import enrich
    enrich(stores.values())
    result = dict(updatedAt=datetime.now(timezone.utc).isoformat(), period=period, cutoff=f'{period}-{max(code_days):02}' if code_days else None, codeLabels=CODE_LABELS, sources=[dict(name=p.name, modifiedAt=datetime.fromtimestamp(p.stat().st_mtime, timezone.utc).isoformat()) for p in [base, attendance]], stores=list(stores.values()), attendance=[dict(day=day, codes=dict(codes), hours=hours[day]) for day, codes in sorted(days.items())], warnings=dict(warnings))
    return result

def publish_data(root):
    files=['site/data.js']
    if (root/'site'/'operacion-data.js').exists():files.append('site/operacion-data.js')
    subprocess.run(['git','add','--',*files],cwd=root,check=True)
    diff=subprocess.run(['git','diff','--cached','--quiet','--',*files],cwd=root)
    if diff.returncode==1:subprocess.run(['git','commit','--only','-m','Actualizar cobertura INVEX','--',*files],cwd=root,check=True)
    elif diff.returncode!=0:raise RuntimeError('No se pudieron comprobar los cambios de datos')
    subprocess.run(['git','push'],cwd=root,check=True)

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base', type=Path, required=True)
    parser.add_argument('--asistencia', type=Path, required=True)
    parser.add_argument('--periodo', required=True, help='AAAA-MM')
    parser.add_argument('--publicar', action='store_true', help='Commit y push de los datos con Git autenticado')
    parser.add_argument('--tablero', type=Path, help='Importar agregados históricos del tablero anterior (opcional)')
    args = parser.parse_args()
    result = generate(args.base, args.asistencia, args.periodo)
    if args.tablero:
        from importar_operacion import import_board
        historical = import_board(args.tablero)
    output = ROOT / 'site' / 'data.js'
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix('.tmp')
    temporary.write_text('window.INVEX_DATA = ' + json.dumps(result, ensure_ascii=False, allow_nan=False) + ';\n', encoding='utf-8')
    temporary.replace(output)
    if args.tablero:
        historical_output = ROOT / 'site' / 'operacion-data.js'
        historical_tmp = historical_output.with_suffix('.tmp')
        historical_tmp.write_text('window.INVEX_OPERACION = '+json.dumps(historical,ensure_ascii=False,allow_nan=False)+';\n',encoding='utf-8')
        historical_tmp.replace(historical_output)
    operating = [s for s in result['stores'] if s['status'] == 'OPERANDO']
    print(json.dumps(dict(tiendas=len(operating), requerida=sum(s['required'] for s in operating), activos=sum(s['active'] for s in operating), vacantes=sum(s['gap'] for s in operating), advertencias=result['warnings']), ensure_ascii=False))
    if args.publicar:
        files = ['site/data.js']
        if (ROOT / 'site' / 'operacion-data.js').exists(): files.append('site/operacion-data.js')
        subprocess.run(['git', 'add', '--', *files], cwd=ROOT, check=True)
        diff = subprocess.run(['git', 'diff', '--cached', '--quiet', '--', *files], cwd=ROOT)
        if diff.returncode == 1:
            subprocess.run(['git', 'commit', '--only', '-m', 'Actualizar cobertura INVEX', '--', *files], cwd=ROOT, check=True)
        elif diff.returncode != 0:
            raise RuntimeError('No se pudieron comprobar los cambios de datos')
        subprocess.run(['git', 'push'], cwd=ROOT, check=True)

if __name__ == '__main__':
    main()
