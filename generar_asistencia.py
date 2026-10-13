"""Equivalente Python del flujo CRUCE.bas incluido en el sistema anterior."""
import unicodedata
from datetime import date,datetime,time
from pathlib import Path
from openpyxl import load_workbook
from engine import to_min

def normalized(value):
    text=unicodedata.normalize('NFKD',str(value or '').strip().upper())
    return ' '.join(''.join(c for c in text if not unicodedata.combining(c)).split())

def name_key(value):return tuple(sorted(normalized(value).split()))

def report_date(value):
    if isinstance(value,datetime):return value.date()
    if isinstance(value,date):return value
    try:return datetime.strptime(str(value).strip(),'%d/%m/%Y').date()
    except ValueError:raise ValueError('El reporte tiene una fecha de fichaje inválida') from None

def minute(value):
    if isinstance(value,datetime):value=value.time()
    if isinstance(value,time):return value.hour*60+value.minute
    if isinstance(value,(float,int)):return round((value%1)*1440) if value%1 else None
    result=to_min(str(value or ''))
    return result if result and 0<result<1440 else None

def summarize(report,day=None):
    workbook=load_workbook(report,read_only=True,data_only=True)
    try:
        sheet=workbook.worksheets[0]
        if sheet.max_column<18 or sheet.max_row<12:raise ValueError('El reporte Pandape no tiene la estructura de fichajes esperada')
        if day is not None:
            dates={}
            for row in sheet.iter_rows(max_row=10,values_only=True):
                values=[v for v in row if v is not None and str(v).strip()]
                for index,value in enumerate(values[:-1]):
                    label=normalized(value).rstrip(':')
                    if label in ('FECHA INICIAL','FECHA FINAL'):
                        raw=values[index+1]
                        parsed=report_date(raw)
                        dates[label]=parsed
            if len(dates)!=2 or not dates['FECHA INICIAL']<=day<=dates['FECHA FINAL']:
                raise ValueError('El rango del reporte no incluye el día que se actualiza')
        headings=next(sheet.iter_rows(min_row=11,max_row=11,values_only=True))
        entries_cols=[i for i,v in enumerate(headings) if normalized(v)=='ENT.']
        exits_cols=[i for i,v in enumerate(headings) if normalized(v)=='SAL.']
        if normalized(headings[2])!='COLABORADOR' or not entries_cols or len(entries_cols)!=len(exits_cols):
            raise ValueError('No se identifican los encabezados Colaborador, Ent. y Sal. del reporte')
        summary={}
        for row in sheet.iter_rows(min_row=12,values_only=True):
            name=row[2];key=name_key(name)
            if not key:continue
            if day is not None:
                row_day=report_date(row[0])
                if row_day!=day:continue
            entries=[m for i in entries_cols if (m:=minute(row[i])) is not None]
            exits=[m for i in exits_cols if (m:=minute(row[i])) is not None]
            data=summary.setdefault(key,{'name':str(name),'entry':None,'exit':None})
            if entries:data['entry']=min(entries+[data['entry']] if data['entry'] is not None else entries)
            if exits:data['exit']=max(exits+[data['exit']] if data['exit'] is not None else exits)
        if not summary:raise ValueError('El reporte Pandape no contiene colaboradores; no se modifica la asistencia')
        return summary
    finally:workbook.close()

def attendance_code(entry,exit,rest,day):
    if entry is None and exit is None:
        weekday=['LUNES','MARTES','MIERCOLES','JUEVES','VIERNES','SABADO','DOMINGO'][day.weekday()]
        text=normalized(rest)
        return 'D' if text==weekday or weekday[:3] in text else 'F'
    # Regla literal de la macro recibida: minutos de entrada del 05 al 15.
    return 'R' if entry is not None and 5<=entry%60<=15 else 'A'

def generate(source,report,destination,day):
    summary=summarize(report,day)
    workbook=load_workbook(source)
    try:
        sheet=workbook['Hoja1']
        headers={normalized(c.value):c.column for c in sheet[1] if c.value is not None}
        required=['NOMBRE COMPLETO','ESTATUS','DESCANSO OFICIAL']
        if any(k not in headers for k in required):raise ValueError('Faltan encabezados obligatorios en la asistencia')
        columns=[c.column for c in sheet[1] if str(c.value).strip()==str(day.day)]
        if len(columns)!=1:raise ValueError('No se identifica una columna única para el día de asistencia')
        column=columns[0]
        if normalized(sheet.cell(1,column+1).value)!='ENTR' or normalized(sheet.cell(1,column+2).value)!='SAL':raise ValueError('El día no tiene las columnas Entr y Sal esperadas')
        audit_column=headers.get('NOMBRE CRUCE (RESUMEN_SALIDAS)',sheet.max_column+1)
        sheet.cell(1,audit_column,'Nombre Cruce (Resumen_Salidas)')
        matched=unmatched=protected=0
        for row in range(2,sheet.max_row+1):
            if normalized(sheet.cell(row,headers['ESTATUS']).value) not in ('ACTIVO','PENDIENTE'):continue
            if normalized(sheet.cell(row,column).value) in ('V','I'):protected+=1;continue
            key=name_key(sheet.cell(row,headers['NOMBRE COMPLETO']).value)
            if not key:continue
            data=summary.get(key)
            if not data:
                sheet.cell(row,audit_column,'Sin coincidencia');sheet.cell(row,column,'no encontrado')
                sheet.cell(row,column+1).value=None;sheet.cell(row,column+2).value=None
                unmatched+=1;continue
            sheet.cell(row,audit_column,data['name'])
            sheet.cell(row,column,attendance_code(data['entry'],data['exit'],sheet.cell(row,headers['DESCANSO OFICIAL']).value,day))
            for offset,field in [(1,'entry'),(2,'exit')]:
                cell=sheet.cell(row,column+offset);cell.value=data[field]/1440 if data[field] is not None else 0;cell.number_format='hh:mm'
            matched+=1
        if not matched:raise ValueError('Ningún colaborador coincide entre Pandape y la plantilla; se conserva la asistencia anterior')
        destination=Path(destination);destination.parent.mkdir(parents=True,exist_ok=True)
        temporary=destination.with_name(destination.stem+'.tmp.xlsx')
        workbook.save(temporary)
        # Abrir el resultado antes de reemplazar la última versión válida.
        check=load_workbook(temporary,read_only=True);check.close();temporary.replace(destination)
        return dict(reportRows=len(summary),matched=matched,unmatched=unmatched,protected=protected,date=day.isoformat())
    finally:workbook.close()
